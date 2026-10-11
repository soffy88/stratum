"""End to end: fixture material -> deck.json (mocked model) -> real .pptx via luobo-ppt.

No network: the model is scripted and illustrations are procedural. Needs a CJK
font for luobo-ppt's text measuring (CI installs fonts-noto-cjk).
"""

from __future__ import annotations

import pytest

from stratum.services.ppt_export.deck_builder import DeckOptions, build_deck
from stratum.services.ppt_export.engine import (
    EngineError,
    apply_qa_fixes,
    parse_outline,
    render_deck,
)
from stratum.services.ppt_export.material import is_language
from stratum.services.ppt_export.styles import list_styles, luobo_home

from .fixtures import FakeProvider, tiny_book


def _has_cjk_font() -> bool:
    import glob
    import os

    pats = [
        "/usr/share/fonts/**/*CJK*",
        "/usr/share/fonts/**/*[Hh]ei*",
        "/usr/share/fonts/**/wqy*",
        os.path.expanduser("~/.fonts/**/*"),
        os.path.expanduser("~/.local/share/fonts/**/*"),
        "/System/Library/Fonts/PingFang*",
    ]
    return any(glob.glob(p, recursive=True) for p in pats)


needs_font = pytest.mark.skipif(not _has_cjk_font(), reason="no CJK font installed")


def test_vendored_skill_is_complete():
    home = luobo_home()
    assert (home / "LICENSE").is_file()
    assert (home / "assets" / "LICENSE-lucide.txt").is_file()
    assert "b6dcb0f325b3ea38a7ca3cd3c211bbba7fb436b2" in (home / "UPSTREAM.md").read_text()
    styles = list_styles()
    assert len(styles) == 14
    assert {"paper-notes", "business-blue", "dark-tech"} <= {s["id"] for s in styles}
    assert all(any("\u4e00" <= ch <= "\u9fff" for ch in s["name"]) for s in styles)


def test_unknown_style_is_rejected(tmp_path):
    with pytest.raises(EngineError):
        render_deck({"meta": {}, "slides": []}, "no-such-style", tmp_path)


def test_parse_outline_lines():
    out = "[WARN] 第 3 页：rows_side 装不下：items 5 条；可改用 list_icons\nnoise\n[ERROR] 第 1 页：缺少标题"
    assert parse_outline(out) == [
        ("WARN", 3, "rows_side 装不下：items 5 条；可改用 list_icons"),
        ("ERROR", 1, "缺少标题"),
    ]


def test_qa_overflow_moves_slide_to_roomier_layout():
    deck = {"slides": [{"n": 2, "type": "grid", "items": [{"text": "a"}] * 4}]}
    qa = {"issues": [{"level": "ERROR", "slide": 2, "check": "overflow", "msg": "x"}]}
    assert apply_qa_fixes(deck, qa) == ["slide 2: layout -> list_icons"]
    assert deck["slides"][0]["layout"] == "list_icons"


@needs_font
@pytest.mark.timeout(300)
@pytest.mark.parametrize("style", ["paper-notes", "dark-tech"])
def test_fixture_book_renders_to_editable_pptx(tmp_path, style):
    from pptx import Presentation

    deck, _ = build_deck(tiny_book(), DeckOptions(pages=8), FakeProvider())
    result = render_deck(deck, style, tmp_path / "work", renderer="lite")

    assert result.pptx.is_file() and result.pptx.stat().st_size > 20_000
    assert result.qa_errors == 0, result.qa.get("issues")
    assert len(result.previews) == 8

    prs = Presentation(str(result.pptx))
    assert len(prs.slides) == 8
    texts, notes = [], []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text)
        if slide.has_notes_slide:
            notes.append(slide.notes_slide.notes_text_frame.text)
    body = " ".join(texts)
    # every title is a real, editable text box (not baked into a picture)
    for s in deck["slides"]:
        title = s.get("title") or s.get("headline")
        if title and s["type"] not in ("toc",):
            assert title.replace("**", "") in body.replace("\n", ""), title
    assert is_language(body, "zh")
    assert sum(1 for n in notes if n.strip()) >= 6
